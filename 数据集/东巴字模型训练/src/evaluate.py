"""Evaluate best validation-selected checkpoint, never use test to select it."""

import argparse
import csv
import json
import time
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
import torch
from common import HOME, atomic_json, code_fingerprint, read_json, sha_file, verified_manifest
from model import GlyphCNN
from PIL import Image
from torch.utils.data import DataLoader
from train import run_lock, set_seed

from data import GlyphDataset, tensor_image


def csv_report(path, fields, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def evaluate(splits, run, output):
    splits, run, output = Path(splits), Path(run), Path(output)
    with run_lock(run):
        return evaluate_locked(splits, run, output)


def evaluate_locked(splits, run, output):
    manifest = verified_manifest(splits)
    # These local checkpoints are trusted pickle files, not uploaded user models.
    last = torch.load(run / "last.pt", map_location="cpu", weights_only=False)
    if not last["progress"]["complete"]:
        raise ValueError("Training is incomplete; do not inspect the held-out test yet")
    best = torch.load(run / "best.pt", map_location="cpu", weights_only=False)
    identity = best["identity"]
    if (
        identity != last["identity"]
        or identity["split_fingerprint"] != manifest["split_fingerprint"]
    ):
        raise ValueError("Checkpoint and split identities differ")
    if identity["code_fingerprint"] != code_fingerprint():
        raise ValueError("Evaluation code differs from the training snapshot")
    config, labels = identity["config"], best["labels"]
    if labels != read_json(splits / "labels.json"):
        raise ValueError("Label mapping changed")
    set_seed(config["seed"], config["threads"])
    model = GlyphCNN(len(labels))
    model.load_state_dict(best["model"])
    model.eval()
    output.mkdir(parents=True, exist_ok=True)
    dataset = GlyphDataset(manifest["dataset_root"], splits / "test.csv")
    confusion = Counter()
    per_total, per_top1, per_top5 = Counter(), Counter(), Counter()
    started = time.perf_counter()
    offset = 0
    with (output / "test_predictions.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "path",
                "expected_id",
                "expected_name",
                "top1_id",
                "top1_name",
                "top5_ids",
                "top1_correct",
                "top5_correct",
                "reference_score",
            ]
        )
        with torch.inference_mode():
            for images, targets in DataLoader(dataset, batch_size=config["batch_size"]):
                probabilities = model(images).softmax(dim=1)
                scores, candidates = probabilities.topk(min(5, len(labels)), dim=1)
                for i, target in enumerate(targets.tolist()):
                    ids, values = candidates[i].tolist(), scores[i].tolist()
                    hit1, hit5 = ids[0] == target, target in ids
                    per_total[target] += 1
                    per_top1[target] += hit1
                    per_top5[target] += hit5
                    if not hit1:
                        confusion[(target, ids[0])] += 1
                    writer.writerow(
                        [
                            dataset.rows[offset]["path"],
                            labels[target]["class_id"],
                            labels[target]["name"],
                            labels[ids[0]]["class_id"],
                            labels[ids[0]]["name"],
                            ";".join(labels[x]["class_id"] for x in ids),
                            hit1,
                            hit5,
                            values[0],
                        ]
                    )
                    offset += 1
    test_seconds = time.perf_counter() - started
    class_rows = [
        {
            "class_id": x["class_id"],
            "name": x["name"],
            "count": per_total[i],
            "top1": per_top1[i] / per_total[i] if per_total[i] else None,
            "top5": per_top5[i] / per_total[i] if per_total[i] else None,
        }
        for i, x in enumerate(labels)
    ]
    confusion_rows = [
        {
            "expected_id": labels[a]["class_id"],
            "expected_name": labels[a]["name"],
            "predicted_id": labels[b]["class_id"],
            "predicted_name": labels[b]["name"],
            "count": count,
        }
        for (a, b), count in confusion.most_common()
    ]
    csv_report(output / "per_class.csv", ["class_id", "name", "count", "top1", "top5"], class_rows)
    csv_report(
        output / "confusions.csv",
        ["expected_id", "expected_name", "predicted_id", "predicted_name", "count"],
        confusion_rows,
    )
    meanings = defaultdict(set)
    for i, item in enumerate(labels):
        for name in [item["name"], *item.get("aliases", [])]:
            meanings[name.strip()].add(i)
    historical_rows = []
    for name, digest in manifest["historical"].items():
        path = Path(manifest["historical_directory"]) / name
        if sha_file(path) != digest:
            raise ValueError("Historical comparison set changed after audit")
        expected = meanings.get(path.stem, set())
        with Image.open(path) as image:
            inputs = tensor_image(image).unsqueeze(0)
        with torch.inference_mode():
            values, candidates = model(inputs).softmax(dim=1).topk(min(5, len(labels)), dim=1)
        ids, scores = candidates[0].tolist(), values[0].tolist()
        valid = len(expected) == 1
        historical_rows.append(
            {
                "file": name,
                "expected_name": path.stem,
                "mapping": "unique" if valid else "ambiguous" if expected else "unmapped",
                "expected_ids": ";".join(labels[x]["class_id"] for x in sorted(expected)),
                "top1_id": labels[ids[0]]["class_id"],
                "top1_name": labels[ids[0]]["name"],
                "top5_names": ";".join(labels[x]["name"] for x in ids),
                "top5_ids": ";".join(labels[x]["class_id"] for x in ids),
                "top1_correct": ids[0] in expected if valid else None,
                "top5_correct": bool(set(ids) & expected) if valid else None,
                "reference_score": scores[0],
            }
        )
    csv_report(output / "historical_comparison.csv", list(historical_rows[0]), historical_rows)
    latency = []
    sample = dataset[0][0].unsqueeze(0)
    with torch.inference_mode():
        for i in range(110):
            tick = time.perf_counter()
            model(sample)
            if i >= 10:
                latency.append((time.perf_counter() - tick) * 1000)
    mapped = [row for row in historical_rows if row["mapping"] == "unique"]
    report = {
        "checkpoint_sha256": sha_file(run / "best.pt"),
        "identity": identity,
        "best_epoch": best["progress"]["epoch"],
        "completed_epochs": last["progress"]["epoch"],
        "test": {
            "count": len(dataset),
            "top1": sum(per_top1.values()) / len(dataset),
            "top5": sum(per_top5.values()) / len(dataset),
            "macro_top1": float(np.mean([x["top1"] for x in class_rows if x["count"]])),
            "macro_top5": float(np.mean([x["top5"] for x in class_rows if x["count"]])),
            "evaluation_seconds": test_seconds,
            "writer_independent": False,
        },
        "historical_comparison": {
            "total": len(historical_rows),
            "unambiguous_labels": len(mapped),
            "top1_correct": sum(bool(x["top1_correct"]) for x in mapped),
            "top5_correct": sum(bool(x["top5_correct"]) for x in mapped),
            "top1": sum(bool(x["top1_correct"]) for x in mapped) / len(mapped) if mapped else None,
            "top5": sum(bool(x["top5_correct"]) for x in mapped) / len(mapped) if mapped else None,
        },
        "forward_only_cpu_ms": {
            "median": float(np.median(latency)),
            "p95": float(np.percentile(latency, 95)),
        },
        "parameters": sum(x.numel() for x in model.parameters()),
        "limitations": [
            "No verified writer identities; exact duplicates removed "
            "but approximate duplicates may remain.",
            "Historical 50-image set has been repeatedly used before; not a fresh blind test.",
            "Softmax values are uncalibrated reference scores, not correctness probabilities.",
            "No unknown rejection or smartphone-photo accuracy demonstrated.",
            "CPU forward timing excludes decode, preprocessing, network and application overhead.",
        ],
    }
    inference_path = output / "inference.pt"
    torch.save(
        {"model": model.state_dict(), "labels": labels, "identity": identity}, inference_path
    )
    report["inference_file_bytes"] = inference_path.stat().st_size
    report["inference_sha256"] = sha_file(inference_path)
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as ExcelImage

    workbook = Workbook()
    summary = workbook.active
    summary.title = "评测说明"
    summary.append(["项目", "结果"])
    for key, value in report.items():
        summary.append([key, json.dumps(value, ensure_ascii=False)])
    for title, rows in [
        ("历史50图对比", historical_rows),
        ("逐类指标", class_rows),
        ("混淆统计", confusion_rows),
    ]:
        sheet = workbook.create_sheet(title)
        if rows:
            sheet.append(list(rows[0]))
            for row in rows:
                sheet.append(list(row.values()))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        if title == "历史50图对比":
            image_column = len(historical_rows[0]) + 1
            sheet.cell(1, image_column, "原图（历史比较集，非盲测）")
            for index, row in enumerate(historical_rows, 2):
                picture = ExcelImage(Path(manifest["historical_directory"]) / row["file"])
                ratio = min(96 / picture.width, 64 / picture.height)
                picture.width, picture.height = picture.width * ratio, picture.height * ratio
                sheet.add_image(picture, sheet.cell(index, image_column).coordinate)
                sheet.row_dimensions[index].height = 52
    workbook.save(output / "东巴字模型评测.xlsx")
    atomic_json(output / "summary.json", report)
    atomic_json(
        run / "status.json",
        {
            "phase": "complete",
            "reports": str(output.resolve()),
            "updated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        },
    )
    print(json.dumps(report, ensure_ascii=False), flush=True)
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--splits", type=Path, default=HOME / "splits")
    parser.add_argument("--run", type=Path, default=HOME / "checkpoints/cpu_v2")
    parser.add_argument("--output", type=Path, default=HOME / "reports/cpu_v2")
    args = parser.parse_args()
    evaluate(args.splits, args.run, args.output)
